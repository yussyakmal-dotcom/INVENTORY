"""Feature parity regressions: money, physical counts, metadata, transfers and exports."""
import base64
import json
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from test_foundation import client, activate, make_product, receipt, central, cmd, main, inv, m
import inventory_restore as r


def ok(response):
    assert response.status_code == 200, response.text
    return response.json()


def setup_stock(client):
    cat=activate(client); wh=central(cat)
    p=make_product(client,uoms={'sak':'25'},purchase_uom='sak',ref_price='10000',min_central='5')
    receipt(client,wh,p['id'],qty='100',cost='10000')
    row=client.get('/api/manufacturing/stock').json()[0]
    return cat,wh,p,row


def supplier(client):
    return ok(cmd(client,'save-supplier',{'code':'SUP','name':'Supplier','lead_days':7}))


def test_po_discount_units_number_and_multiple_batches(client):
    cat,wh,p,row=setup_stock(client); s=supplier(client)
    o=ok(cmd(client,'save-po',{'document_date':str(inv.today()),'warehouse_id':wh,'supplier_id':s['id'],'note':'Catatan utuh',
        'lines':[{'product_id':p['id'],'purchase_uom':'sak','qty':'2','unit_price':'250000','discount':'10'}]}))
    assert o['number'].startswith('PO-') and Decimal(o['data']['total'])==450000
    line=o['data']['lines'][0]
    assert Decimal(line['unit_price'])==9000 and Decimal(line['line_total'])==450000
    o=ok(cmd(client,'submit-po',{'id':o['id'],'revision':o['revision']}))
    o=ok(cmd(client,'approve-operation',{'id':o['id'],'revision':o['revision'],'reason':'Setuju'}))
    o=ok(cmd(client,'receive-po',{'id':o['id'],'revision':o['revision'],'reference':'SJ1','receiver_name':'Petugas',
        'received_date':str(inv.today()),'lines':[
        {'product_id':p['id'],'qty':1,'uom':'sak','batch_number':'A','expiry_date':str(inv.today()+timedelta(days=60)),'rack':'R1'},
        {'product_id':p['id'],'qty':25,'uom':'kg','batch_number':'B','expiry_date':str(inv.today()+timedelta(days=60)),'rack':'R2'}]}))
    assert o['status']=='received'
    rows=client.get('/api/manufacturing/stock').json()
    assert {x['rack'] for x in rows if x['batch_number'] in ['A','B']}=={'R1','R2'}
    assert all(Decimal(x['unit_cost'])==9000 for x in rows if x['batch_number'] in ['A','B'])
    assert o['data']['note']=='Catatan utuh'


def test_blind_count_autosave_excel_recount_and_zero(client):
    _,wh,p,row=setup_stock(client)
    o=ok(cmd(client,'start-opname',{'warehouse_id':wh,'document_date':str(inv.today())}))
    assert 'system_qty' not in o['data']['lines'][0]
    template=client.get(f"/api/manufacturing/opname/{o['id']}/template")
    wb=load_workbook(BytesIO(template.content)); assert wb['_kunci'].sheet_state=='hidden'
    assert 'Saldo' not in str(list(wb['Hitung'].values))
    wb['Hitung']['G2']=0; buf=BytesIO();wb.save(buf)
    imported=ok(cmd(client,'import-count',{'id':o['id'],'revision':o['revision'],'file':base64.b64encode(buf.getvalue()).decode()}))
    assert imported['imported']==1 and imported['errors']==[]
    o=imported['document'];assert o['data']['lines'][0]['counted_qty']=='0'
    o=ok(cmd(client,'submit-opname',{'id':o['id'],'revision':o['revision'],'counted_by':'A','reason':'Nol fisik',
        'lines':[{'balance_id':row['balance_id'],'counted_qty':'0'}]}))
    assert Decimal(o['data']['lines'][0]['difference_value'])==-1000000
    o=ok(cmd(client,'recount-opname',{'id':o['id'],'revision':o['revision'],'reason':'Verifikasi ulang'}))
    assert o['data']['recount']==1 and o['data']['lines'][0]['counted_qty'] is None
    assert 'rounds' not in o['data'] and 'system_qty' not in o['data']['lines'][0]
    assert cmd(client,'import-count',{'id':o['id'],'revision':o['revision'],'file':base64.b64encode(buf.getvalue()).decode()}).status_code==409
    assert cmd(client,'submit-opname',{'id':o['id'],'revision':o['revision'],'counted_by':'A','reason':'Belum',
        'lines':[{'balance_id':row['balance_id'],'counted_qty':None}]}).status_code==422


def test_rack_scopes_and_concurrent_movement_negative_guard(client):
    _,wh,p,row=setup_stock(client)
    ok(cmd(client,'set-rack',{'balance_id':row['balance_id'],'rack':'A'}))
    receipt(client,wh,p['id'],qty=5,batch='OTHER')
    other=next(x for x in client.get('/api/manufacturing/stock').json() if x['batch_number']=='OTHER')
    ok(cmd(client,'set-rack',{'balance_id':other['balance_id'],'rack':'B'}))
    a=ok(cmd(client,'start-opname',{'warehouse_id':wh,'document_date':str(inv.today()),'rack':'A'}))
    b=ok(cmd(client,'start-opname',{'warehouse_id':wh,'document_date':str(inv.today()),'rack':'B'}))
    assert len(a['data']['lines'])==len(b['data']['lines'])==1
    assert cmd(client,'start-opname',{'warehouse_id':wh,'document_date':str(inv.today())}).status_code==409
    ok(cmd(client,'usage',{'warehouse_id':wh,'product_id':p['id'],'qty':99,'reference':'Selama hitung'}))
    a=ok(cmd(client,'submit-opname',{'id':a['id'],'revision':a['revision'],'counted_by':'A','reason':'Fisik',
        'lines':[{'balance_id':row['balance_id'],'counted_qty':'1'}]}))
    assert cmd(client,'approve-operation',{'id':a['id'],'revision':a['revision'],'reason':'Uji konflik'}).status_code==409
    assert Decimal(next(x for x in client.get('/api/manufacturing/stock').json() if x['balance_id']==row['balance_id'])['qty'])==1


def test_draft_reservations_dispatch_cancel_and_transit_loss(client):
    cat,wh,p,row=setup_stock(client);dest=cat['kitchens'][0]['warehouse_id']
    body={'source_id':wh,'destination_id':dest,'shipped_date':str(inv.today()),'lines':[{'product_id':p['id'],'qty':2,'uom':'sak'}]}
    draft=ok(cmd(client,'draft-transfer',body)); assert draft['status']=='draft'
    assert Decimal(client.get('/api/manufacturing/stock').json()[0]['available_qty'])==50
    assert cmd(client,'usage',{'warehouse_id':wh,'product_id':p['id'],'qty':51,'reference':'Reserved'}).status_code==409
    t=ok(cmd(client,'dispatch-transfer',{'id':draft['id'],'revision':draft['revision'],'sender_name':'A','vehicle_number':'BE 123','shipped_date':str(inv.today())}))
    assert Decimal(client.get('/api/manufacturing/stock').json()[0]['qty'])==50
    line=t['lines'][0]
    o=ok(cmd(client,'request-transit-scrap',{'transfer_id':t['id'],'reason':'Hilang dalam perjalanan','lines':[{'id':line['id'],'qty':2}]}))
    assert cmd(client,'receive-transfer',{'id':t['id'],'lines':[{'id':line['id'],'qty':50}]}).status_code==422
    ok(cmd(client,'receive-transfer',{'id':t['id'],'lines':[{'id':line['id'],'qty':48}]}))
    before=client.get('/api/manufacturing/stock').json()
    ok(cmd(client,'approve-operation',{'id':o['id'],'revision':o['revision'],'reason':'Bukti diperiksa'}))
    assert client.get('/api/manufacturing/stock').json()==before
    t=client.get('/api/manufacturing/transfers').json()[0]
    assert t['status']=='closed' and Decimal(t['lines'][0]['lost_qty'])==2
    assert Decimal(t['lines'][0]['received_qty'])==48
    assert cmd(client,'receive-transfer',{'id':t['id'],'lines':[{'id':line['id'],'qty':1}]}).status_code==409


def test_master_excel_roundtrip_preview_rollback(client):
    cat,wh,p,row=setup_stock(client)
    data=client.get('/api/manufacturing/products.xlsx').content
    wb=load_workbook(BytesIO(data)); wb['Produk']['B2']='Nama diperbaiki';buf=BytesIO();wb.save(buf)
    body={'file':base64.b64encode(buf.getvalue()).decode(),'preview':True}
    result=ok(cmd(client,'import-products',body));assert result['errors']==[]
    assert client.get('/api/manufacturing/catalog').json()['products'][0]['name']==p['name']
    ok(cmd(client,'import-products',{**body,'preview':False}))
    updated=client.get('/api/manufacturing/catalog').json()['products'][0]
    assert updated['name']=='Nama diperbaiki' and updated['purchase_uom']=='sak'
    assert Decimal(updated['ref_price'])==10000


def test_document_receipt_atomic_and_reports_scope(client):
    cat=activate(client);p=make_product(client);wh=central(cat);s=supplier(client)
    body={'warehouse_id':wh,'supplier_id':s['id'],'document_date':str(inv.today()),'reference':'SJ','receiver_name':'A',
          'lines':[{'product_id':p['id'],'qty':5,'unit_cost':100,'batch_number':'X','expiry_date':str(inv.today()+timedelta(days=30)),'rack':'R1'}]}
    bad={**body,'lines':body['lines']+[dict(body['lines'][0])]}
    assert cmd(client,'receive-document',bad).status_code==422
    assert client.get('/api/manufacturing/stock').json()==[]
    o=ok(cmd(client,'receive-document',body));assert o['number'].startswith('GRN-')
    report=client.get('/api/manufacturing/report/valuation').json()
    assert Decimal(report[0]['Nilai'])==500
    xlsx=client.get('/api/manufacturing/report/valuation?export=true')
    assert load_workbook(BytesIO(xlsx.content)).active['A2'].value=='Gudang Central'
    with main.SessionLocal() as db:
        u=db.query(main.User).filter_by(username='admin').one();u.role='gudang';u.warehouse_id=None;db.commit()
    assert client.get('/api/manufacturing/report/valuation').json()==[]
    assert client.get(f'/api/manufacturing/report/valuation?warehouse_id={wh}').status_code==403
    assert cmd(client,'save-settings',{'company':'No'}).status_code==403


def test_unified_warehouse_production_assignment_and_capability_guard(client):
    cat=activate(client)
    w=ok(cmd(client,'save-warehouse',{'code':'FACTORY','name':'Gudang produksi','kind':'store','can_produce':True,'can_pack':True}))
    assert w['can_produce'] and w['can_pack']
    with main.SessionLocal() as db:
        assert db.query(m.Kitchen).filter_by(warehouse_id=w['id']).count()==0
    p=make_product(client,'FINISHED',kind='finished',shelf_life_days=3,production_warehouse_id=w['id'])
    assert p['production_warehouse_id']==w['id'] and p['kitchen_id'] is None
    assert cmd(client,'save-warehouse',{**w,'can_produce':False}).status_code==409
    assert cmd(client,'save-warehouse',{**w,'active':False}).status_code==409
    assert cmd(client,'delete-warehouse',w).status_code==409
    material=make_product(client,'MAT')
    recipe=ok(cmd(client,'save-bom',{'product_id':p['id'],'output_qty':40,'materials':[{'product_id':material['id'],'qty':1,'uom':'kg'}],'byproducts':[],'overheads':[]}))
    assert cmd(client,'activate-bom',{'id':recipe['id']}).status_code==200
    assert cmd(client,'save-product',{**material,'production_warehouse_id':w['id']}).status_code==422
    nonprod=next(x for x in cat['warehouses'] if x['kind']=='central')
    assert cmd(client,'save-product',{**p,'production_warehouse_id':nonprod['id'],'change_reason':'Pindah'}).status_code==422


def test_unified_location_upgrade_preserves_balances_and_is_idempotent(client):
    cat,wh,p,row=setup_stock(client)
    kitchen=cat['kitchens'][0]
    finished=make_product(client,'OLD-FINISHED',kind='finished',shelf_life_days=2,kitchen_id=kitchen['id'])
    before=client.get('/api/manufacturing/stock').json()
    with main.SessionLocal() as db:
        db.query(m.ProductionLocation).delete()
        db.get(m.Control,1).schema_version=8
        db.commit()
    inv.initialize(main.engine,main.SessionLocal)
    inv.initialize(main.engine,main.SessionLocal)
    assert client.get('/api/manufacturing/stock').json()==before
    products=client.get('/api/manufacturing/catalog').json()['products']
    assert next(x for x in products if x['id']==finished['id'])['production_warehouse_id']==kitchen['warehouse_id']
    with main.SessionLocal() as db:
        assert db.query(m.ProductionLocation).count()==1
