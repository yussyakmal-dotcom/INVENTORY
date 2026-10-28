"""Restored inventory capabilities, sharing the existing transactional ledger."""
import json
import secrets
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

from sqlalchemy.orm import object_session
import inventory as c
from inventory_models import (Profile, Classification, Product, Supplier, Warehouse, Lot,
                              Balance, Operation, Movement, Audit, Transfer, now)

DEFAULTS = dict(company='INVENTORY', expWarnDays=30, priceTolerance=20,
                safetyDays=7, coverDays=30, usageDays=30, theme='green', density='normal', reserveHours=24)


def profile(db, key):
    row = db.get(Profile, key)
    return json.loads(row.document) if row else {}


def put_profile(db, key, data):
    row = db.get(Profile, key)
    if not row:
        row = Profile(key=key, revision=0)
        db.add(row)
    row.document = c.dump(data)
    row.revision += 1
    db.flush()
    return data


def attached(item, prefix):
    db = object_session(item)
    return profile(db, f'{prefix}:{item.id}') if db and item.id else {}


def settings(db):
    return {**DEFAULTS, **profile(db, 'settings')}


def next_number(db, prefix):
    stamp = c.today().strftime('%y%m')
    key = f'number:{prefix}:{stamp}'
    n = profile(db, key).get('value', 0)
    while True:
        n += 1
        number = f'{prefix}-{stamp}-{n:04d}'
        if not db.query(Operation).filter_by(number=number).first():
            break
    put_profile(db, key, {'value': n})
    return number


def class_dict(row):
    return {k: getattr(row, k) for k in ('id', 'kind', 'name', 'parent_id', 'active', 'revision')} | json.loads(row.document)


def save_classification(db, user, body):
    c.admin(user)
    row = c.get(db, Classification, body['id']) if body.get('id') else Classification(revision=0)
    if row.id and c.integer(body.get('revision'), 'Versi') != row.revision:
        c.fail('Master berubah; muat ulang.', 409)
    kind = body.get('kind')
    if kind not in {'category', 'brand', 'type'}:
        c.fail('Jenis master tidak valid.')
    if row.id and kind != row.kind:
        c.fail('Jenis master tidak dapat diganti.')
    parent_id = c.integer(body['parent_id'], 'Induk') if body.get('parent_id') else None
    seen, depth, parent = {row.id}, 1, parent_id
    while parent:
        if parent in seen or kind != 'category':
            c.fail('Induk kategori tidak valid.')
        seen.add(parent)
        obj = c.get(db, Classification, parent)
        if obj.kind != 'category':
            c.fail('Induk harus kategori.')
        depth += 1
        if depth > 3:
            c.fail('Kategori maksimal tiga tingkat.')
        parent = obj.parent_id
    # Moving a category cannot silently push descendants beyond level three.
    def height(ident):
        children = db.query(Classification).filter_by(parent_id=ident).all()
        return 1 + max((height(x.id) for x in children), default=0)
    if row.id and depth + height(row.id) - 1 > 3:
        c.fail('Pemindahan membuat kategori melebihi tiga tingkat.')
    name = c.label(body.get('name'), 'Nama', 150)
    if db.query(Classification).filter(Classification.kind == kind, Classification.parent_id == parent_id,
            Classification.name == name, Classification.id != (row.id or 0)).first():
        c.fail('Nama sudah digunakan pada kelompok ini.', 409)
    row.kind, row.name, row.parent_id = kind, name, parent_id
    row.active = c.flag(body.get('active', True), 'Aktif')
    row.document = c.dump({k: c.integer(body[k], k, 0) if body.get(k) not in (None, '') else None
                          for k in ('expWarnDays', 'reserveHours')})
    row.revision += 1
    db.add(row)
    db.flush()
    return class_dict(row)


def save_product_profile(db, p, body):
    old = profile(db, f'product:{p.id}')
    data = {**old}
    for key, kind in [('category_id', 'category'), ('brand_id', 'brand'), ('type_id', 'type')]:
        value = body.get(key, old.get(key))
        data[key] = c.integer(value, key) if value not in (None, '') else None
        if data[key] and c.get(db, Classification, data[key]).kind != kind:
            c.fail('Klasifikasi produk tidak sesuai.')
    sid = body.get('supplier_id', old.get('supplier_id'))
    data['supplier_id'] = c.get(db, Supplier, sid).id if sid not in (None, '') else None
    for key in ('min_central', 'min_branch', 'ref_price'):
        data[key] = str(c.number(body.get(key, old.get(key, 0)), key, scale=8 if key == 'ref_price' else 6))
    data['purchase_uom'] = c.selected_unit(p, body.get('purchase_uom') or old.get('purchase_uom') or p.base_uom)[0]
    ttl = body.get('reserve_hours', old.get('reserve_hours'))
    data['reserve_hours'] = c.integer(ttl, 'Batas reservasi', 0) if ttl not in ('', None) else None
    put_profile(db, f'product:{p.id}', data)


def rack(db, balance_id):
    return profile(db, f'balance:{balance_id}').get('rack', '')


def set_rack(db, balance_id, value):
    if not isinstance(value, str) or len(value) > 50:
        c.fail('Rak maksimal 50 karakter.')
    put_profile(db, f'balance:{balance_id}', {'rack': value.strip()})


def restore_archive(db, source):
    """Idempotent backfill only: never overwrite later edits or recreate balances."""
    if db.get(Profile, 'restored-archive'):
        return
    maps = {}
    for group, kind in [('categories', 'category'), ('brands', 'brand'), ('types', 'type')]:
        maps[group] = {}
        pending = list(source.get(group, []))
        while pending:
            progressed = False
            for old in list(pending):
                parent = old.get('parentId')
                if parent and parent not in maps[group]:
                    continue
                row = Classification(kind=kind, name=old['name'], parent_id=maps[group].get(parent),
                    active=old.get('isActive', True), document=c.dump({'expWarnDays': old.get('expiryDays', old.get('expWarnDays')),
                    'reserveHours': old.get('reserveHours', old.get('ttlHours'))}))
                db.add(row); db.flush()
                maps[group][old['id']] = row.id
                pending.remove(old); progressed = True
            if not progressed:
                c.fail('Hierarki kategori arsip tidak valid; data sumber tetap tersimpan.', 409)
    supplier_map = {}
    for old in source.get('suppliers', []):
        s = db.query(Supplier).filter_by(code=old['code']).first()
        if not s:
            s = Supplier(code=old['code'], name=old['name'], contact=old.get('phone', ''), active=old.get('isActive', True))
            db.add(s); db.flush()
            put_profile(db, f'supplier:{s.id}', {'lead_days': old.get('leadTimeDays', old.get('leadDays', 3))})
        supplier_map[old['id']] = s.id
    for old in source.get('products', []):
        p = db.query(Product).filter_by(legacy_id=old['id']).first()
        if p and not db.get(Profile, f'product:{p.id}'):
            put_profile(db, f'product:{p.id}', dict(category_id=maps['categories'].get(old.get('categoryId')),
                brand_id=maps['brands'].get(old.get('brandId')), type_id=maps['types'].get(old.get('typeId')),
                supplier_id=supplier_map.get(old.get('supplierId')), purchase_uom=old.get('poUom') or p.base_uom,
                min_central=str(old.get('minCentral', 0)), min_branch=str(old.get('minBranch', 0)),
                ref_price=str(old.get('refPrice', 0)), reserve_hours=old.get('overrideTtl')))
    for p in db.query(Product).filter(Product.legacy_id.isnot(None)):
        key = f'product:{p.id}'
        meta = profile(db, key)
        prices = {}
        for po in sorted(source.get('pos', []), key=lambda o:o.get('createdAt') or 0, reverse=True):
            if po.get('status') not in {'dipesan','selesai'}: continue
            for line in po.get('lines', []):
                if line.get('productId') == p.legacy_id and line.get('uom') not in prices:
                    prices[line.get('uom', p.base_uom)] = {'value': str(line['price']), 'source': 'Harga terakhir '+po.get('no','')}
        meta['legacy_prices'] = prices
        meta['legacy_lead_days'] = [h['days'] for h in source.get('leadHistory',[]) if h.get('productId')==p.legacy_id and supplier_map.get(h.get('supplierId'))==meta.get('supplier_id')]
        put_profile(db,key,meta)
    for old in source.get('batches', []):
        lot = db.query(Lot).filter_by(legacy_id=old['id']).first()
        if lot:
            for balance in db.query(Balance).filter_by(lot_id=lot.id):
                if not db.get(Profile, f'balance:{balance.id}'):
                    set_rack(db, balance.id, old.get('rack', ''))
    if not db.get(Profile, 'settings'):
        old_settings = source.get('settings', {})
        put_profile(db, 'settings', {**DEFAULTS, **old_settings, 'expWarnDays': old_settings.get('expiryDays', DEFAULTS['expWarnDays']), 'usageDays': old_settings.get('forecastWindow', DEFAULTS['usageDays'])})
    put_profile(db, 'restored-archive', {'at': str(now())})


def public_operation(o, user):
    from inventory_operations import operation_dict
    result = operation_dict(o)
    if o.kind == 'opname' and (o.status == 'draft' or (o.status == 'pending' and user.role not in {'admin', 'approver'})):
        result['data'].pop('rounds', None)
        for line in result['data']['lines']:
            for key in ('system_qty', 'difference', 'unit_cost', 'difference_value'):
                line.pop(key, None)
    return result


def count_save(db, user, body):
    from inventory_operations import role, scope, version, save
    role(user, {'admin', 'gudang'})
    o = c.get(db, Operation, body.get('id'))
    scope(db, user, o.warehouse_id); version(o, body)
    if o.kind != 'opname' or o.status != 'draft':
        c.fail('Sesi tidak sedang menghitung.', 409)
    data = json.loads(o.document)
    if body.get('token') != data.get('token'):
        c.fail('Kunci sesi sudah berubah; unduh template terbaru.', 409)
    by_id = {l['balance_id']: l for l in data['lines']}
    rows = body.get('lines', [])
    if not isinstance(rows, list) or len(rows) > len(by_id):
        c.fail('Baris hitung tidak valid.')
    seen = set()
    for entry in rows:
        bid = c.integer(entry.get('balance_id'), 'Baris')
        if bid not in by_id or bid in seen:
            c.fail('Baris bukan milik sesi atau berulang.')
        seen.add(bid)
        value = entry.get('counted_qty')
        qty = None if value in ('', None) else c.number(value, 'Qty fisik')
        p = c.get(db, Product, by_id[bid]['product_id'])
        if qty is not None and p.discrete and qty != qty.to_integral_value():
            c.fail('Qty produk diskret harus bulat.')
        by_id[bid]['counted_qty'] = None if qty is None else str(qty)
    for key, limit in [('counted_by', 150), ('reason', 500)]:
        if key in body:
            value = body[key]
            if not isinstance(value, str) or len(value) > limit:
                c.fail('Catatan/petugas tidak valid.')
            data[key] = value.strip()
    data['saved_at'] = str(now())
    save(db, o, data)
    return public_operation(o, user)


def recount(db, user, body):
    from inventory_operations import role, scope, version, save, balance_line
    role(user, {'admin', 'approver'})
    o = c.get(db, Operation, body.get('id'))
    scope(db, user, o.warehouse_id); version(o, body)
    if o.kind != 'opname' or o.status != 'pending':
        c.fail('Hanya opname menunggu persetujuan dapat dihitung ulang.', 409)
    data = json.loads(o.document)
    data.setdefault('rounds', []).append({'lines': data['lines'], 'new_lots': data.get('new_lots', []),
        'reason': c.label(body.get('reason'), 'Alasan hitung ulang', 500), 'at': str(now()), 'actor_id': user.id})
    data['lines'] = [{**balance_line(db, c.get(db, Balance, l['balance_id'])), 'counted_qty': None} for l in data['lines']]
    data['new_lots'], data['token'], data['recount'] = [], secrets.token_hex(24), data.get('recount', 0) + 1
    data['reason'], data['counted_by'] = '', ''
    o.status, o.approved_by, o.decided_at = 'draft', None, None
    save(db, o, data)
    return public_operation(o, user)


def xlsx_bytes(sheets, hidden=()):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name[:31])
        for row in rows:
            # Untrusted strings must stay text even if they start with '='.
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str): cell.data_type = 's'
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='176448')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(45, max(14, max(len(str(x.value or '')) for x in column) + 2))
        if name in hidden: ws.sheet_state = 'hidden'
    buf = BytesIO(); wb.save(buf)
    return buf.getvalue()


def opname_template(db, user, ident):
    from inventory_operations import scope
    o = c.get(db, Operation, ident); scope(db, user, o.warehouse_id)
    if o.kind != 'opname' or o.status != 'draft': c.fail('Sesi tidak sedang menghitung.', 409)
    data = json.loads(o.document)
    rows = [['KodeBaris', 'Rak', 'KodeProduk', 'NamaProduk', 'NoBatch', 'Satuan', 'QtyFisik']]
    rows += [[str(l['balance_id']), l.get('rack', ''), l['code'], l['name'], l['batch_number'], l['uom'], l.get('counted_qty')] for l in data['lines']]
    return xlsx_bytes([('Hitung', rows), ('Petunjuk', [['Petunjuk'], ['Isi QtyFisik; isi 0 bila barang tidak ditemukan.'],
        ['Jangan mengubah kode baris, produk, batch, satuan, atau kunci sesi.']]),
        ('_kunci', [['sesi', o.number], ['token', data['token']], ['gudang', str(o.warehouse_id)]])], ['_kunci'])


def opname_import(db, user, body):
    import base64
    from openpyxl import load_workbook
    from inventory_operations import scope
    try:
        raw = base64.b64decode(body.get('file', ''), validate=True)
        if len(raw) > 3_000_000: c.fail('File maksimal 3 MB.')
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        keys = dict(wb['_kunci'].values)
        rows = list(wb['Hitung'].values)
    except (ValueError, KeyError, TypeError, Exception) as exc:
        if hasattr(exc, 'status_code'): raise
        c.fail('File Excel atau kunci sesi tidak valid.')
    o = c.get(db, Operation, body.get('id')); scope(db, user, o.warehouse_id)
    data = json.loads(o.document)
    if keys.get('sesi') != o.number or keys.get('token') != data.get('token') or str(keys.get('gudang')) != str(o.warehouse_id):
        c.fail('Excel bukan milik sesi/gudang ini atau berasal dari hitung ulang sebelumnya.', 409)
    known = {str(l['balance_id']): l for l in data['lines']}
    errors, entries, seen = [], [], set()
    for idx, row in enumerate(rows[1:], 2):
        if not any(x is not None for x in row): continue
        if len(row) < 7: errors.append(f'Baris {idx}: kolom tidak lengkap.'); continue
        key, _, code, _, batch, unit, qty = row[:7]
        line = known.get(str(key))
        if not line or str(key) in seen or str(code) != line['code'] or str(batch) != line['batch_number'] or str(unit) != line['uom']:
            errors.append(f'Baris {idx}: identitas tidak cocok/duplikat.'); continue
        seen.add(str(key))
        if qty is None: continue
        try:
            value = c.number(qty, 'Qty fisik')
            p = c.get(db, Product, line['product_id'])
            if p.discrete and value != value.to_integral_value(): c.fail('Qty harus bulat.')
            entries.append({'balance_id': line['balance_id'], 'counted_qty': str(value)})
        except Exception as exc:
            if not hasattr(exc, 'status_code'): raise
            errors.append(f'Baris {idx}: {exc.detail}')
    result = count_save(db, user, {**body, 'token': data['token'], 'lines': entries})
    return {'document': result, 'imported': len(entries), 'errors': errors}


def receive_document(db, user, body):
    from inventory_operations import role, new_operation, save, snapshot_product
    role(user, {'admin', 'gudang'})
    wh = c.warehouse_access(db, user, body.get('warehouse_id'))
    supplier = c.get(db, Supplier, body.get('supplier_id'))
    if not supplier.active: c.fail('Supplier tidak aktif.')
    entries = body.get('lines')
    if not isinstance(entries, list) or not 1 <= len(entries) <= 200: c.fail('Isi 1–200 baris penerimaan.')
    ref = c.label(body.get('reference'), 'Surat jalan supplier', 100)
    receiver = c.label(body.get('receiver_name'), 'Petugas penerima', 150)
    data = {'supplier_name': supplier.name, 'supplier_code': supplier.code, 'reference': ref,
            'receiver_name': receiver, 'lines': [], 'total': '0'}
    o = new_operation(db, user, 'receipt', body, data, status='posted'); o.supplier_id = supplier.id
    total, seen = Decimal(0), set()
    for row in entries:
        p = c.get(db, Product, row.get('product_id'))
        batch = c.label(row.get('batch_number'), 'Batch', 100)
        if (p.id, batch) in seen: c.fail('Produk dan batch berulang dalam dokumen.')
        seen.add((p.id, batch))
        r = c.receive_stock(db, user, {**row, 'warehouse_id': wh.id, 'reference': o.number})
        line = {**snapshot_product(p), **r, 'batch_number': batch, 'expiry_date': row['expiry_date'], 'rack': row.get('rack', '')}
        data['lines'].append(line); total += Decimal(r['qty']) * Decimal(r['unit_cost'])
    data['total'] = str(total)
    return save(db, o, data)


def restore_command(db, user, operation, body):
    if operation == 'save-classification': return save_classification(db, user, body)
    if operation == 'save-count': return count_save(db, user, body)
    if operation == 'recount-opname': return recount(db, user, body)
    if operation == 'import-count': return opname_import(db, user, body)
    if operation == 'receive-document': return receive_document(db, user, body)
    if operation == 'save-settings':
        c.admin(user)
        data = settings(db)
        data['company'] = c.label(body.get('company'), 'Nama perusahaan', 200)
        for key in ('expWarnDays', 'priceTolerance', 'safetyDays', 'coverDays', 'usageDays', 'reserveHours'):
            data[key] = c.integer(body.get(key, data[key]), key, 1 if key == 'usageDays' else 0)
        if body.get('theme', data['theme']) not in {'green', 'blue', 'brown'}: c.fail('Tema tidak valid.')
        if body.get('density', data['density']) not in {'normal', 'compact'}: c.fail('Kepadatan tidak valid.')
        data['theme'], data['density'] = body.get('theme', data['theme']), body.get('density', data['density'])
        return put_profile(db, 'settings', data)
    if operation == 'set-rack':
        b = c.get(db, Balance, body.get('balance_id')); c.warehouse_access(db, user, b.warehouse_id)
        set_rack(db, b.id, body.get('rack', '')); return {'balance_id': b.id, 'rack': rack(db, b.id)}
    if operation == 'record-print':
        from inventory_operations import scope
        kind = body.get('kind')
        if kind == 'po':
            o = c.get(db, Operation, body.get('id')); scope(db, user, o.warehouse_id)
            if o.kind != 'po' or o.status not in {'approved', 'partial', 'received', 'closed'}: c.fail('PO belum disetujui.')
        elif kind == 'do':
            o = c.get(db, Transfer, body.get('id'))
            ids = c.allowed_warehouses(db, user)
            if ids is not None and not ({o.source_id, o.destination_id} & ids): c.fail('Di luar lingkup.', 403)
        else: c.fail('Jenis cetak tidak valid.')
        key = f'print:{kind}:{o.id}'
        return put_profile(db, key, {'count': profile(db, key).get('count', 0) + 1, 'at': str(now()), 'actor_id': user.id})
    c.fail('Perintah pemulihan tidak ditemukan.', 404)


COMMANDS = {'save-classification', 'save-count', 'recount-opname', 'import-count', 'receive-document',
            'save-settings', 'set-rack', 'record-print'}


def report_rows(db, user, kind, warehouse_id=None, product_id=None, days=30):
    from inventory_operations import listing
    ids = c.allowed_warehouses(db, user)
    if warehouse_id:
        c.warehouse_access(db, user, warehouse_id, write=False)
    rows = []
    if kind in {'valuation', 'expiry'}:
        for r in c.stock_rows(db, user):
            if warehouse_id and r['warehouse_id'] != warehouse_id: continue
            if product_id and r['product_id'] != product_id: continue
            p = c.get(db, Product, r['product_id']); w = c.get(db, Warehouse, r['warehouse_id'])
            rows.append({'Gudang': w.name, 'Kode': p.code, 'Produk': p.name, 'Batch': r['batch_number'],
                'Rak': r['rack'], 'ED': r['expiry_date'], 'Qty': r['qty'], 'Satuan': p.base_uom,
                'HPP': r['unit_cost'], 'Nilai': str(Decimal(r['qty']) * Decimal(r['unit_cost'])),
                'Tersedia': r['available_qty'], 'Status': 'Layak' if r['usable'] else 'Diblokir/kedaluwarsa'})
        return rows
    if kind in {'card', 'usage', 'anomaly'}:
        query = db.query(Movement, Balance, Lot, Product, Warehouse).join(Balance, Movement.balance_id == Balance.id).join(
            Lot, Balance.lot_id == Lot.id).join(Product, Lot.product_id == Product.id).join(Warehouse, Balance.warehouse_id == Warehouse.id)
        query = query.filter(Movement.created_at >= now() - timedelta(days=days))
        if ids is not None: query = query.filter(Balance.warehouse_id.in_(ids))
        if warehouse_id: query = query.filter(Balance.warehouse_id == warehouse_id)
        if product_id: query = query.filter(Product.id == product_id)
        for m, b, lot, p, w in query.order_by(Movement.id.desc()):
            if kind == 'usage' and m.kind != 'USAGE': continue
            if kind == 'anomaly':
                ref = Decimal(profile(db, f'product:{p.id}').get('ref_price', 0))
                if m.kind != 'RECEIPT' or not ref: continue
                deviation = (lot.unit_cost - ref) / ref * 100
                if abs(deviation) <= settings(db)['priceTolerance']: continue
            row = {'Waktu': str(m.created_at), 'Gudang': w.name, 'Kode': p.code, 'Produk': p.name,
                   'Jenis': m.kind, 'Batch': lot.batch_number, 'Qty': str(m.qty_delta), 'Satuan': p.base_uom,
                   'HPP': str(lot.unit_cost), 'Nilai': str(m.qty_delta * lot.unit_cost), 'Dokumen': m.reference, 'Akun': m.actor_id}
            if kind == 'anomaly': row.update(Referensi=str(ref), SelisihPersen=str(deviation))
            rows.append(row)
        if kind == 'anomaly':
            for o in listing(db, user)['documents']:
                if o['kind'] != 'po' or o['status'] in {'cancelled', 'rejected'}: continue
                if warehouse_id and o['warehouse_id'] != warehouse_id: continue
                if o['document_date'] < str(c.today() - timedelta(days=days)): continue
                for l in o['data']['lines']:
                    if product_id and l['product_id'] != product_id: continue
                    ref = Decimal(profile(db, f"product:{l['product_id']}").get('ref_price', 0))
                    if not ref: continue
                    dev = (Decimal(l['unit_price']) - ref) / ref * 100
                    if abs(dev) > settings(db)['priceTolerance']:
                        rows.append({'Waktu': o['document_date'], 'Gudang': c.get(db, Warehouse, o['warehouse_id']).name,
                            'Kode': l['code'], 'Produk': l['name'], 'Jenis': 'PO', 'Batch': '', 'Qty': l['qty'], 'Satuan': l['uom'],
                            'HPP': l['unit_price'], 'Nilai': l.get('line_total', str(Decimal(l['qty']) * Decimal(l['unit_price']))),
                            'Dokumen': o['number'], 'Akun': o['created_by'], 'Referensi': str(ref), 'SelisihPersen': str(dev)})
        return rows
    if kind == 'scrap':
        for o in listing(db, user)['documents']:
            if o['kind'] != 'scrap' or (warehouse_id and o['warehouse_id'] != warehouse_id): continue
            if o['document_date'] < str(c.today() - timedelta(days=days)): continue
            for l in o['data']['lines']:
                if product_id and l['product_id'] != product_id: continue
                rows.append({'Tanggal': o['document_date'], 'Dokumen': o['number'], 'Gudang': c.get(db, Warehouse, o['warehouse_id']).name,
                    'Produk': l['name'], 'Batch': l['batch_number'], 'Qty': l['qty'], 'Satuan': l['uom'],
                    'Nilai': str(Decimal(l['qty']) * Decimal(l['unit_cost'])), 'Sumber': o['data'].get('source', 'manual'),
                    'Alasan': o['data']['reason'], 'Status': o['status']})
        return rows
    c.fail('Jenis laporan tidak tersedia.')


def forecast(db, user):
    from inventory_operations import listing
    cfg = settings(db); days = cfg['usageDays']
    stock = c.stock_rows(db, user)
    usage = report_rows(db, user, 'usage', days=days)
    docs = listing(db, user)['documents']
    rows = []
    for p in db.query(Product).filter_by(active=True):
        meta = profile(db, f'product:{p.id}')
        daily = sum(-Decimal(x['Qty']) for x in usage if x['Kode'] == p.code) / days
        onhand = sum(Decimal(x['available_qty']) for x in stock if x['product_id'] == p.id and x['usable'])
        onorder = sum(Decimal(l['qty']) - Decimal(l['received_qty']) for o in docs if o['kind'] == 'po' and o['status'] in {'approved', 'partial'} for l in o['data']['lines'] if l['product_id'] == p.id)
        supplier = db.get(Supplier, meta.get('supplier_id')) if meta.get('supplier_id') else None
        lead = attached(supplier, 'supplier').get('lead_days', 3) if supplier else 3
        actual = list(meta.get('legacy_lead_days', []))
        for o in docs:
            if o['kind'] != 'po' or o['supplier_id'] != (supplier.id if supplier else None): continue
            if not any(l['product_id'] == p.id for l in o['data']['lines']): continue
            actual.extend((datetime.fromisoformat(r['received_date']).date() - datetime.fromisoformat(o['document_date']).date()).days for r in o['data'].get('receipts', []))
        if actual: lead = sum(actual) / len(actual)
        lead_d = Decimal(str(lead))
        rop = daily * (lead_d + cfg['safetyDays'])
        target = daily * (lead_d + cfg['safetyDays'] + cfg['coverDays'])
        target = max(target, Decimal(meta.get('min_central', 0)))
        unit = meta.get('purchase_uom', p.base_uom)
        factor = Decimal(({p.base_uom: '1', **json.loads(p.uoms)}).get(unit, 1))
        shortage = max(Decimal(0), target - onhand - onorder)
        from decimal import ROUND_CEILING
        qty = (shortage / factor).to_integral_value(rounding=ROUND_CEILING)
        rows.append({'product_id': p.id, 'code': p.code, 'name': p.name, 'supplier_id': supplier.id if supplier else None,
            'supplier_name': supplier.name if supplier else 'Belum diatur', 'daily': str(daily), 'lead_days': str(lead),
            'onhand': str(onhand), 'onorder': str(onorder), 'rop': str(rop), 'target': str(target),
            'qty': str(qty), 'purchase_uom': unit, 'base_uom': p.base_uom,
            'status': 'Pesan' if onhand + onorder <= rop and qty else 'Pantau' if qty else 'Aman'})
    return rows


def install(app, ready_user, get_db, user_model):
    from fastapi import Depends, Query
    from fastapi.responses import Response

    @app.get('/api/manufacturing/inventory-tools')
    def inventory_tools(u=Depends(ready_user), db=Depends(get_db)):
        c.require_active(db)
        return {'classifications': [class_dict(x) for x in db.query(Classification).order_by(Classification.name)],
                'settings': settings(db)}

    @app.get('/api/manufacturing/opname/{ident}/template')
    def count_template(ident: int, u=Depends(ready_user), db=Depends(get_db)):
        return Response(opname_template(db, u, ident), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="opname-{ident}.xlsx"'})

    @app.get('/api/manufacturing/report/{kind}')
    def report(kind: str, warehouse_id: int = None, product_id: int = None, days: int = Query(30, ge=1, le=3660),
               export: bool = False, u=Depends(ready_user), db=Depends(get_db)):
        c.require_active(db)
        rows = report_rows(db, u, kind, warehouse_id, product_id, days)
        if export:
            heads = list(rows[0]) if rows else ['Tidak ada data']
            return Response(xlsx_bytes([('Laporan', [heads] + [[r.get(h) for h in heads] for r in rows])]),
                media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                headers={'Content-Disposition': f'attachment; filename="laporan-{kind}.xlsx"'})
        return rows

    @app.get('/api/manufacturing/forecast')
    def forecast_list(u=Depends(ready_user), db=Depends(get_db)):
        c.require_active(db)
        return forecast(db, u)

    @app.get('/api/manufacturing/audit')
    def audit_list(after: int = 0, limit: int = Query(200, ge=1, le=1000), u=Depends(ready_user), db=Depends(get_db)):
        c.admin(u)
        q = db.query(Audit).filter(Audit.id > after).order_by(Audit.id).limit(limit)
        names = {x.id: x.name for x in db.query(user_model)}
        return [{'id': x.id, 'at': str(x.created_at), 'actor': names.get(x.actor_id, 'Sistem'),
                 'operation': x.operation, 'data': json.loads(x.document)} for x in q]

PRODUCT_HEADERS = ['kode', 'nama', 'jenis', 'merek', 'kategori', 'satuan_dasar', 'konversi', 'satuan_beli',
                   'min_central', 'min_cabang', 'harga_referensi', 'supplier', 'ttl_jam', 'status',
                   'jenis_manufaktur', 'umur_simpan', 'dapur', 'harga_jual', 'diskret']


def category_path(db, ident):
    parts, seen = [], set()
    while ident and ident not in seen:
        seen.add(ident); row = db.get(Classification, ident)
        if not row: break
        parts.insert(0, row.name); ident = row.parent_id
    return ' > '.join(parts)


def products_excel(db, template=False):
    rows = [PRODUCT_HEADERS]
    if not template:
        for p in db.query(Product).order_by(Product.code):
            d = profile(db, f'product:{p.id}')
            def name(key):
                row = db.get(Classification, d[key]) if d.get(key) else None
                return row.name if row else ''
            supplier = db.get(Supplier, d['supplier_id']) if d.get('supplier_id') else None
            from inventory_models import Kitchen
            kitchen = db.get(Kitchen, p.kitchen_id) if p.kitchen_id else None
            rows.append([p.code, p.name, name('type_id'), name('brand_id'), category_path(db, d.get('category_id')),
                p.base_uom, '; '.join(f'{k}={v}' for k,v in json.loads(p.uoms).items()), d.get('purchase_uom', p.base_uom),
                d.get('min_central', '0'), d.get('min_branch', '0'), d.get('ref_price', '0'), supplier.code if supplier else '',
                d.get('reserve_hours'), 'aktif' if p.active else 'nonaktif', p.kind, p.shelf_life_days,
                kitchen.code if kitchen else '', str(p.selling_price), 'ya' if p.discrete else 'tidak'])
    return xlsx_bytes([('Produk', rows), ('Petunjuk', [['Petunjuk'],
        ['Kode unik; kode yang sudah ada diperbarui. Kolom tambahan manufaktur kosong mempertahankan atribut lama saat update.'],
        ['Konversi contoh sak=25; dus=10. Kategori contoh Bahan > Tepung > Terigu.'],
        ['Jenis manufaktur: raw, packaging, wip, finished, package, byproduct. Produk jadi wajib dapur dan umur simpan.'],
        ['Supplier menggunakan kode master. Preview wajib diperiksa sebelum simpan. Tidak ada perubahan bila satu baris gagal.']])])


def import_products(db, user, body):
    import base64
    from openpyxl import load_workbook
    from inventory_models import Kitchen
    c.admin(user)
    try:
        raw = base64.b64decode(body.get('file', ''), validate=True)
        if len(raw) > 3_000_000: c.fail('File maksimal 3 MB.')
        from itertools import islice
        if raw.startswith(b'PK'):
            wb = load_workbook(BytesIO(raw), read_only=True, data_only=False)
            rows = list(islice(wb['Produk'].values, 2002))
        else:
            import csv
            rows = list(islice(csv.reader(raw.decode('utf-8-sig').splitlines()), 2002))
    except Exception as exc:
        if hasattr(exc, 'status_code'): raise
        c.fail('File produk Excel tidak valid.')
    if not rows or list(rows[0])[:6] != PRODUCT_HEADERS[:6]: c.fail('Header tidak sesuai template produk.')
    if len(rows) > 2001: c.fail('Maksimal 2000 produk per impor.')
    results, errors, seen = [], [], set()
    outer = db.begin_nested()
    def classification(kind, path):
        parent = None
        names = str(path).split('>') if kind == 'category' else [path]
        if len(names) > 3: c.fail('Kategori maksimal tiga tingkat.')
        for name in names:
            name = c.label(name, 'Nama klasifikasi', 150)
            row = db.query(Classification).filter_by(kind=kind, parent_id=parent, name=name).first()
            if not row:
                row = Classification(kind=kind, name=name, parent_id=parent, active=True)
                db.add(row); db.flush()
            parent = row.id
        return parent
    for idx, row in enumerate(rows[1:], 2):
        if not any(x is not None for x in row): continue
        nested = db.begin_nested()
        try:
            values = dict(zip(PRODUCT_HEADERS, row))
            code = c.label(values.get('kode'), 'Kode produk', 50).upper()
            if code in seen: c.fail('Kode berulang dalam file.')
            seen.add(code)
            p = db.query(Product).filter_by(code=code).first()
            payload = c.product_dict(p) if p else {'kind': 'raw', 'discrete': False, 'active': True, 'shelf_life_days': None, 'kitchen_id': None}
            payload.update(code=code, name=values.get('nama'), base_uom=values.get('satuan_dasar'))
            units = {}
            for pair in str(values.get('konversi') or '').split(';'):
                if not pair.strip(): continue
                parts = pair.split('=')
                if len(parts) != 2: c.fail('Format konversi nama=faktor.')
                units[parts[0].strip()] = parts[1].strip()
            payload['uoms'] = units
            for column, key in [('jenis','type_id'), ('merek','brand_id'), ('kategori','category_id')]:
                if values.get(column): payload[key] = classification({'jenis':'type','merek':'brand','kategori':'category'}[column], values[column])
            for column, key in [('min_central','min_central'),('min_cabang','min_branch'),('harga_referensi','ref_price'),('ttl_jam','reserve_hours'),('satuan_beli','purchase_uom'),('jenis_manufaktur','kind'),('umur_simpan','shelf_life_days'),('harga_jual','selling_price')]:
                if values.get(column) is not None: payload[key] = values[column]
            if values.get('supplier'):
                supplier = db.query(Supplier).filter_by(code=str(values['supplier'])).first()
                if not supplier: c.fail('Kode supplier tidak ditemukan.')
                payload['supplier_id'] = supplier.id
            if values.get('dapur'):
                kitchen = db.query(Kitchen).filter_by(code=str(values['dapur'])).first()
                if not kitchen: c.fail('Kode dapur tidak ditemukan.')
                payload['kitchen_id'] = kitchen.id
                payload['change_reason'] = 'Impor master produk'
            if values.get('status'):
                if values['status'] not in ('aktif','nonaktif'): c.fail('Status harus aktif/nonaktif.')
                payload['active'] = values['status'] == 'aktif'
            if values.get('diskret'):
                if values['diskret'] not in ('ya','tidak'): c.fail('Diskret harus ya/tidak.')
                payload['discrete'] = values['diskret'] == 'ya'
            old_id = p.id if p else None
            result = c.save_product(db, user, payload, old_id)
            results.append({'id': old_id, 'code': result['code'], 'name': result['name']})
            nested.commit()
        except Exception as exc:
            nested.rollback()
            if not hasattr(exc, 'status_code'): raise
            errors.append(f'Baris {idx}: {exc.detail}')
    preview = body.get('preview', True)
    if preview or errors: outer.rollback()
    else: outer.commit()
    if errors and not preview: c.fail('Impor dibatalkan: ' + ' '.join(errors))
    return {'rows': results, 'errors': errors, 'preview': preview}


COMMANDS.add('import-products')
_original_restore_command = restore_command

def restore_command(db, user, operation, body):
    if operation == 'import-products': return import_products(db, user, body)
    return _original_restore_command(db, user, operation, body)

_original_install = install

def install(app, ready_user, get_db, user_model):
    from fastapi import Depends
    from fastapi.responses import Response
    _original_install(app, ready_user, get_db, user_model)

    @app.get('/api/manufacturing/products.xlsx')
    def product_export(template: bool = False, u=Depends(ready_user), db=Depends(get_db)):
        c.require_active(db)
        return Response(products_excel(db, template), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename="' + ('template-produk' if template else 'produk') + '.xlsx"'})


def transfer_ttl(db, p):
    data = profile(db, f'product:{p.id}')
    if data.get('reserve_hours') is not None: return int(data['reserve_hours'])
    ident = data.get('category_id')
    while ident:
        row = c.get(db, Classification, ident); doc = json.loads(row.document)
        if doc.get('reserveHours') is not None: return int(doc['reserveHours'])
        ident = row.parent_id
    return settings(db)['reserveHours']


def transfer_plan(db, user, body):
    from inventory_operations import balance_line
    src = c.warehouse_access(db, user, body.get('source_id'))
    dst = c.get(db, Warehouse, body.get('destination_id'))
    if not dst.active or src.id == dst.id: c.fail('Tujuan harus aktif dan berbeda dari asal.')
    rows = body.get('lines')
    if not isinstance(rows, list) or not 1 <= len(rows) <= 200: c.fail('Isi 1–200 baris pengiriman.')
    lines, seen, ttl = [], set(), []
    for row in rows:
        p = c.get(db, Product, row.get('product_id'))
        if not p.active or p.id in seen: c.fail('Produk harus aktif dan tidak berulang.')
        seen.add(p.id)
        unit, factor = c.selected_unit(p, row.get('uom'))
        qty = c.quantity_for(p, c.number(row.get('qty'), 'Qty', True) * factor)
        for b, lot, take in c.fefo(db, p, src, qty):
            lines.append({**balance_line(db, b), 'qty': str(take)})
        ttl.append(transfer_ttl(db, p))
    return {'lines': lines, 'ttl_hours': min(ttl), 'destination_id': dst.id, 'source_id': src.id,
            'total': str(sum(Decimal(l['qty'])*Decimal(l['unit_cost']) for l in lines))}


def draft_transfer(db, user, body):
    from inventory_operations import new_operation, operation_dict
    plan = transfer_plan(db, user, body)
    # Identity is needed on dispatch; a draft may be prepared before a driver is assigned.
    data = {**plan, 'dispatch': body, 'expires_at': str(now()+timedelta(hours=plan['ttl_hours']))}
    o = new_operation(db, user, 'transfer', {'number': body.get('document_number'), 'warehouse_id': plan['source_id'],
        'document_date': body.get('shipped_date') or str(c.today())}, data)
    return operation_dict(o)


def dispatch_draft(db, user, body):
    from inventory_operations import version, save
    o = c.get(db, Operation, body.get('id'))
    c.warehouse_access(db, user, o.warehouse_id); version(o, body)
    data = json.loads(o.document)
    if o.kind != 'transfer' or o.status != 'draft' or datetime.fromisoformat(data['expires_at']) <= now():
        c.fail('Draft tidak aktif atau reservasi kedaluwarsa.', 409)
    payload = {**data['dispatch'], 'document_number': o.number,
               'sender_name': body.get('sender_name'), 'vehicle_number': body.get('vehicle_number'),
               'shipped_date': body.get('shipped_date') or str(c.today())}
    # Release only this reservation, within the same transaction; failures roll it back.
    o.status = 'posted'; db.flush()
    picks = []
    from inventory_operations import reserved
    for l in data['lines']:
        b = c.get(db, Balance, l['balance_id']); lot = c.get(db, Lot, b.lot_id)
        qty = Decimal(l['qty'])
        if lot.blocked or lot.expiry_date is None or lot.expiry_date < c.today() or b.qty-reserved(db, b.id)<qty:
            c.fail('Batch draft tidak lagi layak/tersedia; batalkan dan buat draft baru.', 409)
        picks.append((b,lot,qty))
    result = c.ship(db, user, payload, reserved_picks=picks)
    data['transfer_id'] = result['id']; save(db, o, data)
    return result


def transit_loss(db, line_id, pending=False):
    from inventory_models import TransferLine
    total = Decimal(0)
    for o in db.query(Operation).filter(Operation.kind == 'scrap', Operation.status.in_(['posted','pending'] if pending else ['posted'])):
        data = json.loads(o.document)
        if data.get('source') == 'transit':
            total += sum(Decimal(l['qty']) for l in data['lines'] if l['transfer_line_id']==line_id)
    return total


def request_transit_scrap(db, user, body):
    from inventory_models import TransferLine
    from inventory_operations import new_operation, operation_dict, snapshot_product
    t = c.get(db, Transfer, body.get('transfer_id'))
    c.warehouse_access(db, user, t.destination_id)
    rows=body.get('lines')
    if not isinstance(rows,list) or not 1<=len(rows)<=200: c.fail('Isi baris kehilangan transit.')
    lines, seen = [], set()
    for entry in rows:
        line=c.get(db,TransferLine,entry.get('id'))
        if line.transfer_id!=t.id or line.id in seen: c.fail('Baris transfer salah/berulang.')
        seen.add(line.id); lot=c.get(db,Lot,line.lot_id); p=c.get(db,Product,lot.product_id)
        qty=c.number(entry.get('qty'), 'Qty hilang')
        if not qty: continue
        c.quantity_for(p,qty)
        if qty>line.shipped_qty-line.received_qty-transit_loss(db,line.id,pending=True): c.fail('Qty hilang melebihi sisa transit.',409)
        lines.append({**snapshot_product(p),'transfer_line_id':line.id,'lot_id':lot.id,'batch_number':lot.batch_number,
                      'qty':str(qty),'unit_cost':str(lot.unit_cost)})
    if not lines: c.fail('Isi qty kehilangan lebih dari nol.')
    data={'source':'transit','transfer_id':t.id,'reason':c.label(body.get('reason'),'Alasan kehilangan',500),
          'lines':lines,'total':str(sum(Decimal(l['qty'])*Decimal(l['unit_cost']) for l in lines))}
    o=new_operation(db,user,'scrap',{'number':body.get('number'),'warehouse_id':t.destination_id,
         'document_date':body.get('document_date') or str(c.today())},data,status='pending')
    return operation_dict(o)


def post_transit_scrap(db, o, data):
    from inventory_models import TransferLine
    t=c.get(db,Transfer,data['transfer_id'])
    for l in data['lines']:
        line=c.get(db,TransferLine,l['transfer_line_id'])
        if Decimal(l['qty'])>line.shipped_qty-line.received_qty-transit_loss(db,line.id):
            c.fail('Sisa transit berubah; periksa kembali pengajuan.',409)
    # Stock left the origin on dispatch. Approval records loss, never deducts stock again.
    o.status='posted';db.flush()
    lines=db.query(TransferLine).filter_by(transfer_id=t.id).all()
    if all(l.shipped_qty==l.received_qty+transit_loss(db,l.id) for l in lines): t.status='closed'


COMMANDS.update({'preview-transfer','draft-transfer','dispatch-transfer','request-transit-scrap'})
_product_restore_command = restore_command

def restore_command(db, user, operation, body):
    actions={'preview-transfer':transfer_plan,'draft-transfer':draft_transfer,'dispatch-transfer':dispatch_draft,'request-transit-scrap':request_transit_scrap}
    if operation in actions:return actions[operation](db,user,body)
    return _product_restore_command(db,user,operation,body)


def audit_before(db, operation, payload):
    ident = payload.get('id')
    if operation == 'save-settings': return settings(db)
    if operation == 'set-rack': return profile(db, f"balance:{payload.get('balance_id')}")
    if not ident: return None
    from inventory_operations import operation_dict, supplier_dict
    if operation == 'save-product':
        row = db.get(Product, c.integer(ident, 'ID')); return c.product_dict(row) if row else None
    if operation == 'save-supplier':
        row = db.get(Supplier, c.integer(ident, 'ID')); return supplier_dict(row) if row else None
    if operation == 'save-classification':
        row = db.get(Classification, c.integer(ident, 'ID')); return class_dict(row) if row else None
    if operation in {'receive-transfer'}:
        row = db.get(Transfer, c.integer(ident, 'ID')); return c.transfer_dict(db,row) if row else None
    if operation in {'save-po','submit-po','receive-po','save-count','import-count','recount-opname','submit-opname',
                     'approve-operation','reject-operation','cancel-operation','close-po','dispatch-transfer'}:
        row = db.get(Operation,c.integer(ident,'ID')); return operation_dict(row) if row else None
    return None


def legacy_movements(db, user, kind, warehouse_id, product_id, days):
    """Read archived transactions for reporting; never replay them into balances."""
    ctrl = c.control(db)
    if not ctrl.archive: return []
    archive = json.loads(ctrl.archive)
    whs = {w.legacy_id:w for w in db.query(Warehouse) if w.legacy_id is not None}
    products = {p.legacy_id:p for p in db.query(Product) if p.legacy_id is not None}
    ids = c.allowed_warehouses(db,user)
    since = now()-timedelta(days=days)
    rows=[]
    for m in archive.get('movements',[]):
        w,p=whs.get(m.get('warehouseId')),products.get(m.get('productId'))
        if not w or not p or (ids is not None and w.id not in ids):continue
        if warehouse_id and w.id!=warehouse_id:continue
        if product_id and p.id!=product_id:continue
        at=c.legacy_timestamp(m.get('ts'))
        if at<since:continue
        if kind=='usage' and m.get('type')!='USE':continue
        if kind=='anomaly' and (m.get('type')!='IN' or m.get('ref')=='SALDO-AWAL'):continue
        qty=Decimal(str(m.get('qty',0)));cost=Decimal(str(m.get('cost',0)))
        if kind=='anomaly':
            ref=Decimal(profile(db,f'product:{p.id}').get('ref_price',0))
            if not ref:continue
            dev=(cost-ref)/ref*100
            if abs(dev)<=settings(db)['priceTolerance']:continue
        row={'Waktu':str(at),'Gudang':w.name,'Kode':p.code,'Produk':p.name,'Jenis':m.get('type'),
             'Batch':m.get('batchNo',''),'Qty':str(qty),'Satuan':p.base_uom,'HPP':str(cost),'Nilai':str(qty*cost),
             'Dokumen':m.get('ref',''),'Akun':m.get('user',''),'Sumber':'Sebelum migrasi'}
        if kind=='anomaly':row.update(Referensi=str(ref),SelisihPersen=str(dev))
        rows.append(row)
    return rows

_current_report_rows = report_rows

def report_rows(db, user, kind, warehouse_id=None, product_id=None, days=30):
    rows=_current_report_rows(db,user,kind,warehouse_id,product_id,days)
    if kind in {'card','usage','anomaly'}:
        for row in rows:row['Sumber']='Fondasi'
        rows+=legacy_movements(db,user,kind,warehouse_id,product_id,days)
        rows.sort(key=lambda r:r['Waktu'],reverse=True)
    return rows
